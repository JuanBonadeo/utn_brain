#!/usr/bin/env python3
"""Transfiere filas CSV_PEATONAL de AnyLogic a 04-resultados-corridas.xlsx.

Entradas aceptadas: el archivo que escribe el experimento PeatonalCorridasApareadas
(`corridas_peatonales.csv`) o un texto copiado de la consola de AnyLogic. Solo se leen las
líneas que empiezan con un prefijo CSV_PEATONAL*; el resto se ignora.

Por defecto el script solo valida y muestra un resumen (no escribe). Para cargar:

    python3 cargar_corridas.py corridas_peatonales.csv --escribir

Reglas de seguridad:
- Las filas CSV_PEATONAL_DEMO solo se aceptan con --demo y nunca se escriben sobre la planilla
  oficial: exigen --salida hacia otra copia.
- Una fila CSV_PEATONAL_INCOMPLETO (corrida que no drenó) bloquea la carga.
- Cada semilla debe tener E0 y E1, con igual cantidad de pasajeros generados.
- No se sobrescriben celdas cargadas con otro valor salvo con --sobrescribir.

n de pares variable (T1.5): la planilla no tiene un tamaño fijo de 30 pares. El cargador
acepta semillas contiguas desde SEMILLA_INICIAL hasta la cantidad de filas que la planilla
ofrezca (se generan con `construir_planilla.py`), y reporta cuántos pares quedaron cargados.
El t crítico para el intervalo de confianza se calcula en la propia planilla (hoja Resumen)
con una fórmula T.INV.2T en función del n real y de k (comparaciones primarias), no a mano.
"""
from __future__ import annotations

import argparse
import math
import re
import sys
from pathlib import Path

AQUI = Path(__file__).resolve().parent
PLANILLA_OFICIAL = AQUI / '04-resultados-corridas.xlsx'
HOJA = 'Corridas'
HOJA_RESUMEN = 'Resumen'
FILA_ENCABEZADO = 7
PRIMERA_FILA = 8
SEMILLA_INICIAL = 20260923
COLUMNA_SEMILLA = 2
MIN_PARES = 10       # n mínimo para habilitar el IC en la hoja Resumen (n0, §7.3 del plan)
OBJETIVO_PARES = 30  # tamaño de diseño de referencia (no es un mínimo fijo)

# Orden exacto de los campos 3..17 de filaResultadoPed() y de los bloques de la hoja Corridas.
CAMPOS = [
    'Pasajeros generados',
    'Pasajeros desviados',
    'Procesados a las 09:30',
    'Procesados con drenaje',
    'Espera media',
    'P90 de espera',
    'Proporción con espera mayor a 30 s',
    'Lq',
    'Cola máxima',
    'Tiempo de disipación',
    'Utilización media',
    'Dispersión de utilización',
    'Pasajeros de la cohorte pico',
    'Espera media de la cohorte pico',
    'P90 de la cohorte pico',
]
ENTEROS = {0, 1, 2, 3, 8, 12}
FRACCIONES = {6, 10}
PREFIJO = re.compile(r'^(CSV_PEATONAL(?:_DEMO|_INCOMPLETO)?);(.*)$')


class ErrorCarga(Exception):
    pass


def columnas(indice: int) -> tuple[int, int, int]:
    """Columnas 1-based de E0, E1 y diferencia para el campo `indice`."""
    base = 3 + 3 * indice
    return base, base + 1, base + 2


def leer_filas(rutas: list[Path]) -> tuple[dict[str, list[tuple[str, list[str]]]], int]:
    por_tipo: dict[str, list[tuple[str, list[str]]]] = {}
    ignoradas = 0
    for ruta in rutas:
        for n, linea in enumerate(ruta.read_text(encoding='utf-8', errors='replace').splitlines(), 1):
            m = PREFIJO.match(linea.strip())
            if not m:
                ignoradas += 1
                continue
            por_tipo.setdefault(m.group(1), []).append((f'{ruta.name}:{n}', m.group(2).split(';')))
    return por_tipo, ignoradas


def convertir(origen: str, campos: list[str]) -> tuple[str, int, list[float]]:
    if len(campos) != 2 + len(CAMPOS):
        raise ErrorCarga(f'{origen}: se esperaban {2 + len(CAMPOS)} campos y hay {len(campos)}')
    escenario = campos[0]
    if escenario == 'E1B':
        raise ErrorCarga(f'{origen}: escenario E1B (28 molinetes, ampliación hipotética) no se carga '
                          f'en esta planilla; solo se cargan E0 (20) y E1 (22)')
    if escenario not in ('E0', 'E1'):
        raise ErrorCarga(f'{origen}: escenario {escenario!r}; solo se cargan E0 (20) y E1 (22)')
    try:
        semilla = int(campos[1])
        valores = [float(x) for x in campos[2:]]
    except ValueError as exc:
        raise ErrorCarga(f'{origen}: valor no numérico ({exc})') from None
    for i, v in enumerate(valores):
        if not math.isfinite(v) or v < 0:
            raise ErrorCarga(f'{origen}: {CAMPOS[i]} inválido ({v})')
        if i in ENTEROS and v != int(v):
            raise ErrorCarga(f'{origen}: {CAMPOS[i]} debe ser entero ({v})')
        if i in FRACCIONES and v > 1:
            raise ErrorCarga(f'{origen}: {CAMPOS[i]} debe estar en [0, 1] ({v})')
    generados, desviados, al_corte, drenaje = valores[0:4]
    if desviados != 0:
        raise ErrorCarga(f'{origen}: E0-E1 no desvía pasajeros ({desviados})')
    if drenaje + desviados != generados:
        raise ErrorCarga(f'{origen}: conservación violada (generados={generados}, drenaje={drenaje})')
    if al_corte > drenaje:
        raise ErrorCarga(f'{origen}: procesados al corte mayor que con drenaje')
    if valores[12] > generados:
        raise ErrorCarga(f'{origen}: cohorte pico mayor que el total generado')
    return escenario, semilla, valores


def armar_pares(filas: list[tuple[str, list[str]]]) -> dict[int, dict[str, list[float]]]:
    pares: dict[int, dict[str, list[float]]] = {}
    repetidas = []
    for origen, campos in filas:
        escenario, semilla, valores = convertir(origen, campos)
        previo = pares.setdefault(semilla, {}).get(escenario)
        if previo is not None and previo != valores:
            raise ErrorCarga(f'{origen}: semilla {semilla} {escenario} repetida con valores distintos')
        if previo is not None:
            repetidas.append(origen)
        pares[semilla][escenario] = valores
    if repetidas:
        print(f'Aviso: {len(repetidas)} filas idénticas repetidas se cuentan una sola vez '
              f'({", ".join(repetidas[:3])}{"…" if len(repetidas) > 3 else ""}). '
              'Borrar el archivo de corridas antes de relanzar el experimento.')
    for semilla, par in sorted(pares.items()):
        faltan = {'E0', 'E1'} - set(par)
        if faltan:
            raise ErrorCarga(f'semilla {semilla}: falta {", ".join(sorted(faltan))}; el par está incompleto')
        if par['E0'][0] != par['E1'][0]:
            raise ErrorCarga(f'semilla {semilla}: E0 y E1 generaron distinta demanda '
                             f'({par["E0"][0]:.0f} vs {par["E1"][0]:.0f}); no son corridas apareadas')
    return pares


def verificar_encabezados(hoja) -> None:
    for i, nombre in enumerate(CAMPOS):
        c0, c1, cd = columnas(i)
        esperados = (f'E0 {nombre}', f'E1 {nombre}', f'Diferencia {nombre}')
        reales = tuple(hoja.cell(FILA_ENCABEZADO, c).value for c in (c0, c1, cd))
        if reales != esperados:
            raise ErrorCarga(f'La planilla no coincide con CSV_PEATONAL en el campo {i + 3}: {reales} != {esperados}')
    if hoja.cell(FILA_ENCABEZADO, COLUMNA_SEMILLA).value != 'Semilla':
        raise ErrorCarga('La columna B de la planilla no es Semilla')


def leer_capacidad(hoja) -> dict[int, int]:
    """Semilla -> fila, leyendo desde PRIMERA_FILA mientras haya semillas contiguas.

    Reemplaza el ULTIMA_FILA fijo: la planilla puede tener cualquier cantidad de filas
    (construidas por construir_planilla.py), siempre que las semillas sean contiguas desde
    SEMILLA_INICIAL. Esto es lo que le permite al cargador aceptar n de pares variable.
    """
    fila_por_semilla: dict[int, int] = {}
    fila = PRIMERA_FILA
    esperada = SEMILLA_INICIAL
    while True:
        valor = hoja.cell(fila, COLUMNA_SEMILLA).value
        if valor is None:
            break
        if valor != esperada:
            raise ErrorCarga(f'fila {fila}: semilla {valor!r} no es la esperada {esperada} '
                              f'(deben ser contiguas desde {SEMILLA_INICIAL})')
        fila_por_semilla[valor] = fila
        fila += 1
        esperada += 1
    if not fila_por_semilla:
        raise ErrorCarga('la planilla no tiene semillas cargadas en la columna B')
    return fila_por_semilla


def cargar(pares, planilla: Path, salida: Path, sobrescribir: bool) -> int:
    import openpyxl

    libro = openpyxl.load_workbook(planilla)
    hoja = libro[HOJA]
    verificar_encabezados(hoja)
    fila_por_semilla = leer_capacidad(hoja)
    escritas = 0
    for semilla, par in sorted(pares.items()):
        fila = fila_por_semilla.get(semilla)
        if fila is None:
            raise ErrorCarga(f'semilla {semilla} no figura en la columna Semilla de la planilla')
        for i in range(len(CAMPOS)):
            for escenario, col in zip(('E0', 'E1'), columnas(i)[:2]):
                valor = par[escenario][i]
                valor = int(valor) if i in ENTEROS else valor
                celda = hoja.cell(fila, col)
                if celda.value not in (None, '') and celda.value != valor and not sobrescribir:
                    raise ErrorCarga(f'{celda.coordinate} ya contiene {celda.value}; use --sobrescribir si corresponde')
                celda.value = valor
        escritas += 1
    libro.save(salida)
    return escritas


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('entradas', nargs='+', type=Path, help='archivo corridas_peatonales.csv o texto de consola')
    ap.add_argument('--planilla', type=Path, default=PLANILLA_OFICIAL)
    ap.add_argument('--salida', type=Path, help='destino; por defecto, la misma planilla')
    ap.add_argument('--escribir', action='store_true', help='guardar los valores (sin esto solo valida)')
    ap.add_argument('--demo', action='store_true', help='cargar filas CSV_PEATONAL_DEMO en una copia de prueba')
    ap.add_argument('--sobrescribir', action='store_true')
    args = ap.parse_args(argv)

    try:
        por_tipo, _ = leer_filas(args.entradas)
        if por_tipo.get('CSV_PEATONAL_INCOMPLETO'):
            origenes = ', '.join(o for o, _ in por_tipo['CSV_PEATONAL_INCOMPLETO'])
            raise ErrorCarga(f'hay corridas que no drenaron ({origenes}); repetirlas antes de cargar')
        tipo = 'CSV_PEATONAL_DEMO' if args.demo else 'CSV_PEATONAL'
        otro = 'CSV_PEATONAL' if args.demo else 'CSV_PEATONAL_DEMO'
        if por_tipo.get(otro):
            raise ErrorCarga(f'la entrada mezcla filas {otro}; separar demo y producción')
        filas = por_tipo.get(tipo, [])
        if not filas:
            raise ErrorCarga(f'no se encontraron filas {tipo}')
        pares = armar_pares(filas)
        print(f'{len(filas)} filas {tipo} leídas; {len(pares)} pares E0-E1 completos '
              f'(semillas {min(pares)}..{max(pares)})')
        if len(pares) < MIN_PARES:
            print(f'Aviso: faltan al menos {MIN_PARES - len(pares)} pares para alcanzar el mínimo '
                  f'n = {MIN_PARES} (el IC de la hoja Resumen no se habilita antes).')
        elif len(pares) < OBJETIVO_PARES:
            print(f'Aviso: {OBJETIVO_PARES - len(pares)} pares más completan el diseño de referencia '
                  f'de {OBJETIVO_PARES} (el IC ya es válido con n = {len(pares)}).')
        if not args.escribir:
            print('Validación sin escritura. Agregar --escribir para cargar la planilla.')
            return 0
        salida = (args.salida or args.planilla).resolve()
        if args.demo and salida == PLANILLA_OFICIAL.resolve():
            raise ErrorCarga('las filas de demostración no se escriben en la planilla oficial; usar --salida')
        n = cargar(pares, args.planilla, salida, args.sobrescribir)
        print(f'OK: {n} pares cargados en {salida}')
        return 0
    except ErrorCarga as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
