#!/usr/bin/env python3
"""Verifica estructura y lógica del prototipo; no sustituye Build/Run en AnyLogic.
Requiere AnyLogic 8 PLE instalado en /Applications (macOS).
"""
from pathlib import Path
import xml.etree.ElementTree as E
import subprocess
p=Path(__file__).with_name('SubteConstitucion.alp')
r=E.parse(p).getroot();a=r.find('Model/ActiveObjectClasses/ActiveObjectClass')
fields=[]
for v in a.findall('Variables/Variable'):
 prop=v.find('Properties'); val=prop.findtext('DefaultValue/Code') if v.get('Class')=='Parameter' else prop.findtext('InitialValue/Code')
 fields.append(f'{prop.findtext("Type")} {v.findtext("Name")} = {val};')
functions=[]
for f in a.findall('Functions/Function'):
 args=', '.join(x.findtext('Type')+' '+x.findtext('Name') for x in f.findall('Parameter'))
 functions.append(f'public {f.findtext("ReturnType")} {f.findtext("Name")}({args}) {{\n{f.findtext("Body")}\n}}')
common='\n'.join(fields+functions)
# Compile actual signatures against the installed AnyLogic libraries.
real='''import com.anylogic.engine.*;
import com.anylogic.libraries.processmodeling.*;
public class SubteApiCheck extends Agent {
Source<Agent> source;
EventTimeout proximaTanda, cierreFranja, observarCola;
'''+common+'\n}'
import tempfile
workspace=tempfile.TemporaryDirectory(prefix='subte-check-')
d=Path(workspace.name)
(d/'SubteApiCheck.java').write_text(real)
plugins=Path('/Applications/AnyLogic 8 PLE.app/Contents/Resources/Java/plugins')
jars=list(plugins.glob('com.anylogic*/**/*.jar'))
javac='/Applications/AnyLogic 8 PLE.app/Contents/jre/bin/javac'
subprocess.run([javac,'-cp',':'.join(map(str,jars)),str(d/'SubteApiCheck.java')],check=True)
print('OK: Java functions compile against installed AnyLogic 8.9.9 API')
# Exercise the model functions with explicit event traces, not a second queue model.
harness='''public class SubteLogicCheck {
static class Agent {}
static class Event { void restart(double t) {} }
static class Source { int injected; void inject(int n) {injected+=n;} }
static class Engine { void finish() {} }
Source source=new Source();
Event proximaTanda=new Event(), cierreFranja=new Event(), observarCola=new Event();
double clock;
double time() {return clock;}
Engine getEngine(){return new Engine();}
void traceln(String s){}
'''+common+'''
static void eq(double a,double b){if(Math.abs(a-b)>1e-8)throw new AssertionError(a+" != "+b);}
public static void main(String[] args){
 SubteLogicCheck m=new SubteLogicCheck();m.modoDemo=false;
 try{m.inicializar();throw new AssertionError("Missing field data accepted");}catch(IllegalArgumentException expected){}
 m.modoDemo=true;m.molinetesBase=2;m.usarPerfilSBASE=false;m.inicializar();
 Agent[] p={new Agent(),new Agent(),new Agent(),new Agent(),new Agent()};
 Agent u=new Agent(),v=new Agent();
 m.nPrincipal=5;m.nGenerados=5;m.tandasAbiertas.add(0.0);
 // A single batch: two immediately served, three wait.
 for(Agent x:p)m.entraCola(x);
 m.iniciaServicio(p[0],u);m.iniciaServicio(p[1],v);eq(m.qMax,3);
 m.clock=3;m.liberaMolinete(u);m.salePasajero();m.iniciaServicio(p[2],u);
 m.liberaMolinete(v);m.salePasajero();m.iniciaServicio(p[3],v);
 m.clock=6;m.liberaMolinete(u);m.salePasajero();m.iniciaServicio(p[4],u);
 m.liberaMolinete(v);m.salePasajero();m.observar();
 eq(m.sumaEspera/5,2.4);eq(m.percentil90(),6);eq(m.disipaciones.get(0),6);
 m.clock=9;m.liberaMolinete(u);m.salePasajero();m.observar();
 m.clock=9000;m.cerrarFranja();m.observar();eq(m.areaCola,12);eq(m.areaOcupados,15);
 eq(m.nEsperando,0);eq(m.nOcupados,0);eq(m.pendientesAlCierre,0);
 if(!m.finalizado)throw new AssertionError("Not finished");
 // E2 zero/full diversion, E1 capacity, E3 service and batch cut-off.
 for(int e=0;e<4;e++){
  SubteLogicCheck n=new SubteLogicCheck();n.escenario=e;n.fraccionDesvioE2=1;n.inicializar();n.generarTanda();
  eq(n.nGenerados,n.nPrincipal+n.nDesviados);
  eq(n.source.injected,n.nPrincipal);
  if(e==2)eq(n.nPrincipal,0);
  eq(n.capacidadMolinetes(),e==1?28:20);eq(n.servicioSeg(),e==3?2.5:3);
  int total=n.nGenerados;n.clock=9000;n.generarTanda();eq(n.nGenerados,total);
 }
 // A passenger still in service at 09:30 is counted only in drained throughput.
 SubteLogicCheck z=new SubteLogicCheck();z.inicializar();z.nPrincipal=1;z.nGenerados=1;
 z.clock=8999;Agent late=new Agent();z.entraCola(late);z.iniciaServicio(late,u);
 z.clock=9000;z.cerrarFranja();eq(z.pendientesAlCierre,1);
 z.clock=9002;z.liberaMolinete(u);z.salePasajero();z.observar();
 eq(z.nProcesadosFranja,0);eq(z.nProcesados,1);eq(z.areaOcupados,1);
 System.out.println("OK: missing data, exact batch waits/P90/Lq/occupation, E0-E3, balance and 09:30 drain");
}
}'''
(d/'SubteLogicCheck.java').write_text(harness)
subprocess.run([javac,str(d/'SubteLogicCheck.java')],check=True)
subprocess.run(['/Applications/AnyLogic 8 PLE.app/Contents/jre/bin/java','-cp',str(d),'SubteLogicCheck'],check=True)
# Compile the Java bodies of the incremental pedestrian root with small API stubs.
ped_agent=next(x for x in r.findall('Model/ActiveObjectClasses/ActiveObjectClass') if x.findtext('Name')=='MainPeatonal')
ped_fields=[]
for v in ped_agent.findall('Variables/Variable'):
 prop=v.find('Properties'); val=prop.findtext('DefaultValue/Code') if v.get('Class')=='Parameter' else prop.findtext('InitialValue/Code')
 ped_fields.append(f'{prop.findtext("Type")} {v.findtext("Name")} = {val};')
ped_functions=[]
for f in ped_agent.findall('Functions/Function'):
 args=', '.join(x.findtext('Type')+' '+x.findtext('Name') for x in f.findall('Parameter'))
 ped_functions.append(f'public {f.findtext("ReturnType")} {f.findtext("Name")}({args}) {{\n{f.findtext("Body")}\n}}')
ped_check='''public class SubtePedLogicCheck {
static class Pasajero {
 double tEntradaColaPed, tIngresoSistemaPed, tInicioServicioPed, servicioAsignadoPed;
 boolean enColaPed, enPicoPed;
 String nombreMolinetePed="";
}
static class PedSourceStub { int injected, calls; java.util.ArrayList<Integer> sizes=new java.util.ArrayList<Integer>(); void inject(int n) {injected+=n;calls++;sizes.add(n);} }
static class EventStub { double next=-1; void restart(double t) {next=t;} }
static class ServiceStub { int calls, suspended; void setServiceSuspended(Object p, boolean value) { calls++; if(value)suspended++; } }
static class EngineStub { boolean finished; void finish(){finished=true;} }
static class RandomStub extends java.util.Random { long seed=-1; public synchronized void setSeed(long s){seed=s;super.setSeed(s);} }
PedSourceStub pedSource=new PedSourceStub(); EventStub proximaTandaPed=new EventStub(), cierreMetricasPed=new EventStub();
ServiceStub molinetesPeatonales=new ServiceStub(); EngineStub engine=new EngineStub(); RandomStub rng=new RandomStub();
Object molinetePed21=new Object(), molinetePed22=new Object(), molinetePed23=new Object(), molinetePed24=new Object();
Object molinetePed25=new Object(), molinetePed26=new Object(), molinetePed27=new Object(), molinetePed28=new Object();
java.util.ArrayList<String> trazas=new java.util.ArrayList<String>();
double clock; double time(){return clock;} void traceln(String s){trazas.add(s);}
EngineStub getEngine(){return engine;} java.util.Random getDefaultRandomGenerator(){return rng;}
'''+ '\n'.join(ped_fields+ped_functions) + '''
static void ok(boolean c,String m){if(!c)throw new AssertionError(m);}
static void eq(double a,double b,String m){ok(Math.abs(a-b)<1e-9,m+": "+a+" != "+b);}
public static void main(String[] args) throws Exception {
 SubtePedLogicCheck m=new SubtePedLogicCheck();
 m.configurarMolinetesPed();
 ok(m.molinetesPeatonales.calls==8 && m.molinetesPeatonales.suspended==8,"E0 suspende 8");
 SubtePedLogicCheck e1=new SubtePedLogicCheck();e1.molinetesOperativosPed=28;e1.configurarMolinetesPed();
 ok(e1.molinetesPeatonales.calls==8 && e1.molinetesPeatonales.suspended==0,"E1 habilita 28");
 // Demo: corte de metricas en horizonteMetricasPedSeg y semilla efectiva semillaPed.
 m.inicioPicoPedSeg=0;m.finPicoPedSeg=10;m.inicializarPed();
 eq(m.horizonteCortePed,600,"corte demo");ok(m.rng.seed==20260923L,"semilla efectiva");
 Pasajero p=new Pasajero();m.clock=2;m.registraIngresoPed(p);p.tEntradaColaPed=m.time();m.entraColaPed(p);
 eq(p.servicioAsignadoPed,3,"servicio demo asignado");
 m.clock=5;p.nombreMolinetePed="molinetePed01";m.comienzaServicioPed(p);
 m.clock=8;m.terminaServicioPed(p);
 ok(m.nEsperandoPed==0 && m.nEsperasPed==1 && Math.abs(m.esperaMediaPed()-3)<1e-9,"espera");
 ok(m.percentil90Ped()==3 && m.proporcionMas30Ped()==0 && m.esperasPicoPed.size()==1 && m.p90PicoPed()==3,"P90 y cohorte");
 eq(m.utilizacionMediaPed(),3.0/(600*20),"utilizacion");
 ok(m.utilizacionPorMolinetePed().contains("molinetePed01=0.005"),"utilizacion por puesto");
 m.nProcesadosHorizontePed=1;
 ok(m.filaResultadoPed().startsWith("E0;20260923;1;0;1;"),"fila");
 ok(m.filaResultadoPed().split(";").length==17,"17 campos");
 Pasajero tarde=new Pasajero();tarde.nombreMolinetePed="molinetePed01";tarde.tInicioServicioPed=599;
 m.clock=605;m.terminaServicioPed(tarde);
 eq(m.ocupacionPorMolinetePed.get("molinetePed01"),4,"ocupacion recortada al corte");
 // Demo completa: seis tandas, cierre de arribos, drenaje y prefijo DEMO sin finalizar el motor.
 SubtePedLogicCheck d=new SubtePedLogicCheck();d.inicializarPed();
 for(int k=0;k<6;k++){d.clock=30*k;d.generarTandaPed();}
 ok(d.arribosCerradosPed && d.nTandasPed==6 && d.nInyectadosPed==480 && d.pedSource.injected==480,"demo 6x80");
 d.clock=700;d.generarTandaPed();ok(d.nInyectadosPed==480,"sin arribos despues del cierre");
 for(int k=0;k<480;k++){Pasajero x=new Pasajero();d.registraIngresoPed(x);}
 for(int k=0;k<479;k++)d.salePed();
 ok(!d.resultadoEmitidoPed,"no emite antes de drenar");
 d.salePed();ok(d.resultadoEmitidoPed && !d.engine.finished,"demo emite al drenar");
 ok(d.trazas.get(d.trazas.size()-1).startsWith("CSV_PEATONAL_DEMO;E0;20260923;480;0;"),"prefijo demo");
 d.emitirIncompletoPed();ok(!d.trazas.get(d.trazas.size()-1).startsWith("CSV_PEATONAL_INCOMPLETO"),"sin fila incompleta tras emitir");
 // Franja: sin datos calibrados se bloquea.
 SubtePedLogicCheck f=new SubtePedLogicCheck();f.modoFranjaPed=true;
 try{f.inicializarPed();throw new AssertionError("Franja sin calibrar aceptada");}catch(IllegalArgumentException expected){}
 f.emitirIncompletoPed();ok(f.trazas.isEmpty() && !f.resultadoEmitidoPed,"sin fila incompleta si la inicializacion fallo");
 // Franja con entradas de prueba: perfil SBASE por ventana, cierre a 9000 s y corte de metricas al cierre.
 f.tamanoTandaPed=100;f.intervaloTandaPedSeg=900;f.primeraTandaPedSeg=0;f.servicioPedSeg=2;f.inicializarPed();
 eq(f.horizonteCortePed,9000,"corte franja");
 int esperado=0;for(double v:f.perfilPed15min)esperado+=(int)Math.round(100*v/1748.2);
 for(int k=0;k<10;k++){f.clock=900*k;f.generarTandaPed();ok(f.pedSource.sizes.get(k)==(int)Math.round(100*f.perfilPed15min[k]/1748.2),"tanda "+k);}
 ok(f.arribosCerradosPed && f.nTandasPed==10 && f.nInyectadosPed==esperado,"diez ventanas y cierre 09:30");
 f.clock=9000;f.generarTandaPed();ok(f.nInyectadosPed==esperado,"sin arribos a las 09:30");
 SubtePedLogicCheck g=new SubtePedLogicCheck();g.modoFranjaPed=true;g.tamanoTandaPed=10;g.intervaloTandaPedSeg=1000;
 g.primeraTandaPedSeg=500;g.servicioPedSeg=2;g.usarPerfilPed=false;g.inicializarPed();
 java.io.File salida=java.io.File.createTempFile("corridas",".csv");salida.delete();g.archivoSalidaPed=salida.getPath();
 for(int k=0;k<9;k++){g.clock=500+1000*k;g.generarTandaPed();}
 ok(g.arribosCerradosPed && g.nInyectadosPed==90,"perfil constante y ultimo arribo antes de 9000");
 Pasajero[] q=new Pasajero[90];for(int k=0;k<90;k++){q[k]=new Pasajero();g.clock=4600;g.registraIngresoPed(q[k]);}
 ok(g.esperasPicoPed.isEmpty() && q[0].enPicoPed && q[0].servicioAsignadoPed==2,"cohorte pico por ingreso y servicio franja");
 for(int k=0;k<89;k++)g.salePed();
 g.clock=9000;g.actualizarAreaColaPed();g.nProcesadosHorizontePed=g.nProcesadosPed;ok(!g.drenadoPed(),"pendiente al cierre");
 g.clock=9050;g.salePed();
 ok(g.resultadoEmitidoPed && g.engine.finished,"franja emite y finaliza al drenar");
 java.util.List<String> lineas=java.nio.file.Files.readAllLines(salida.toPath());
 ok(lineas.size()==1 && lineas.get(0).startsWith("CSV_PEATONAL;E0;20260923;90;0;89;90;"),"archivo de salida: "+lineas);
 eq(g.ultimaDisipacionPed,9050-8500,"disipacion desde la ultima tanda");
 // Conservacion y corrida incompleta.
 SubtePedLogicCheck h=new SubtePedLogicCheck();h.inicializarPed();h.nInyectadosPed=2;h.nGeneradosPed=1;h.nProcesadosPed=2;
 try{h.emitirResultadoPed();throw new AssertionError("Conservacion violada aceptada");}catch(IllegalStateException expected){}
 h.emitirIncompletoPed();ok(h.trazas.get(h.trazas.size()-1).startsWith("CSV_PEATONAL_INCOMPLETO;E0;"),"fila incompleta");
 salida.delete();
 System.out.println("OK: pedestrian demo/franja generation, 09:30 cut, drain, conservation, CSV prefixes and output file");
}
}'''
(d/'SubtePedLogicCheck.java').write_text(ped_check)
subprocess.run([javac,str(d/'SubtePedLogicCheck.java')],check=True)
subprocess.run(['/Applications/AnyLogic 8 PLE.app/Contents/jre/bin/java','-cp',str(d),'SubtePedLogicCheck'],check=True)
# Compile the real AnyLogic API call used to suspend the last eight physical service points.
ped_api='''import com.anylogic.engine.*;
import com.anylogic.engine.markup.*;
public class SubtePedApiCheck extends Agent {
 ServiceWLine<ServicePoint<QueuePath>> services;
 ServicePoint<QueuePath> point;
 void configure(){ services.setServiceSuspended(point, true); point.getName(); }
 void seedAndFinish(long semilla){ getDefaultRandomGenerator().setSeed(semilla); getEngine().finish(); }
}'''
(d/'SubtePedApiCheck.java').write_text(ped_api)
subprocess.run([javac,'-cp',':'.join(map(str,jars)),str(d/'SubtePedApiCheck.java')],check=True)
# Parameter-variation entries reference parameter Ids; every other Id must be unique.
references={id(x) for tag in ('FreeformParamValue','RangeVariationParamValue') for v in r.iter(tag) for x in v.findall('Id')}
ids=[x.text for x in r.iter('Id') if id(x) not in references];assert len(ids)==len(set(ids)), 'Duplicate IDs'
return_modifiers={x.text for x in r.iter('ReturnModificator')}
assert return_modifiers <= {'VOID','RETURNS_VALUE'}, f'Invalid function return modifier: {return_modifiers}'
assert len(a.findall('EmbeddedObjects/EmbeddedObject'))==7
experiments=r.findall('Model/Experiments/SimulationExperiment')
for e in [x for x in experiments if x.findtext('Name') in {'E0','E1','E2','E3'}]:
 ps={p.findtext('ParameterName'):p.findtext('ParameterValue/Code') for p in e.findall('Parameters/Parameter')}
 assert ps['escenario']==e.findtext('Name')[1:]
agents={x.findtext('Name'):x for x in r.findall('Model/ActiveObjectClasses/ActiveObjectClass')}
ped=agents['MainPeatonal']
main=agents['Main']
assert main.findtext('CurrentLevel')=='1783514567443'
main_level=main.find('Presentation/Level')
assert main_level is not None and main_level.findtext('Id')=='1783514567443'
assert main_level.find("Presentation/Rectangle[Name='panelTablero']") is not None
for agent, min_x in ((a,1100),(ped,1000)):
 for section in ('Variables','Functions','Events'):
  technical=agent.findall(f'{section}/*')
  assert all(x.findtext('PresentationFlag')=='false' for x in technical)
  assert all(int(x.findtext('X'))>=min_x for x in technical)
ped_blocks={x.findtext('Name'):x.findtext('ActiveObjectClass/ClassName') for x in ped.findall('EmbeddedObjects/EmbeddedObject')}
assert ped_blocks=={'pedSource':'PedSource','pedMolinetes':'PedService','pedSalida':'PedGoTo','pedSink':'PedSink'}
assert agents['Pasajero'] is not None
assert any(e.findtext('Name')=='PeatonalDemo' and e.get('ActiveObjectClassId')==ped.findtext('Id') for e in experiments)
ped_experiments={e.findtext('Name'):e for e in experiments if e.findtext('Name') in {'PeatonalE0','PeatonalE1'}}
assert set(ped_experiments)=={'PeatonalE0','PeatonalE1'}
for name, expected in {'PeatonalE0':'20','PeatonalE1':'28'}.items():
 ps={p.findtext('ParameterName'):p.findtext('ParameterValue/Code') for p in ped_experiments[name].findall('Parameters/Parameter')}
 assert ps['molinetesOperativosPed']==expected
 assert ps['semillaPed']=='20260923L'
 assert ped_experiments[name].findtext('SeedValue')=='20260923'
 assert ped_experiments[name].findtext('ModelTimeProperties/FinalTime')=='900'
callbacks_delay=next(x for x in ped.findall('EmbeddedObjects/EmbeddedObject') if x.findtext('Name')=='pedMolinetes').find("Parameters/Parameter[Name='delayTime']").findtext('Value/Code')
service_points=ped.findall('.//ServicePoint')
assert len(service_points)==28
assert {x.findtext('Name') for x in service_points}=={f'molinetePed{i:02d}' for i in range(1,29)}
assert ped.find(".//TargetLine[Name='salidaPeatonal']") is not None
assert 'configurarMolinetesPed()' in ped.findtext('StartupCode')
startup=ped.findtext('StartupCode')
assert startup.index('inicializarPed()') < startup.index('configurarMolinetesPed()')
assert 'cierreMetricasPed.restart(horizonteCortePed)' in startup
assert 'proximaTandaPed.restart(modoFranjaPed ? primeraTandaPedSeg : 0)' in startup
assert ped.findtext('DestroyCode')=='emitirIncompletoPed();'
assert callbacks_delay=='ped.servicioAsignadoPed'
assert ped.find(".//Text[Name='etiquetaMolinetes']").findtext('TextCode')=='"Molinetes activos: " + molinetesOperativosPed + "/28"'
ped_service=next(x for x in ped.findall('EmbeddedObjects/EmbeddedObject') if x.findtext('Name')=='pedMolinetes')
callbacks={x.findtext('Name'):x.findtext('Value/Code') for x in ped_service.findall('Parameters/Parameter')}
assert callbacks['onBeginService']=='ped.nombreMolinetePed = service.getName(); comienzaServicioPed(ped);'
assert callbacks['onEndService']=='terminaServicioPed(ped);'
ped_variable_names={x.findtext('Name') for x in ped.findall('Variables/Variable')}
assert {'esperasPed','esperasPicoPed','ocupacionPorMolinetePed','nProcesadosHorizontePed','semillaPed','resultadoEmitidoPed'} <= ped_variable_names
ped_events={x.findtext('Name'):x for x in ped.findall('Events/Event')}
assert 'cierreMetricasPed' in ped_events
assert 'nProcesadosHorizontePed = nProcesadosPed' in ped_events['cierreMetricasPed'].findtext('Action')
assert 'if (drenadoPed()) emitirResultadoPed();' in ped_events['cierreMetricasPed'].findtext('Action')
# Franja parametrizada: los datos de campo siguen sin valor y la demo no cambia.
ped_defaults={v.findtext('Name'):v.findtext('Properties/DefaultValue/Code') for v in ped.findall("Variables/Variable[@Class='Parameter']")}
for name in ('tamanoTandaPed','intervaloTandaPedSeg','primeraTandaPedSeg','servicioPedSeg'):
 assert ped_defaults[name]=='-1', f'{name} debe quedar pendiente de calibración'
assert ped_defaults['modoFranjaPed']=='false' and ped_defaults['horizonteArribosPedSeg']=='9000'
main_defaults={v.findtext('Name'):v.findtext('Properties/DefaultValue/Code') for v in a.findall("Variables/Variable[@Class='Parameter']")}
assert ped_defaults['perfilPed15min']==main_defaults['perfil15min']
assert ped_defaults['archivoSalidaPed']=='""'
demo=next(e for e in experiments if e.findtext('Name')=='PeatonalDemo')
assert demo.findtext('ModelTimeProperties/FinalTime')=='900'
ped_param_ids={v.findtext('Id'):v.findtext('Name') for v in ped.findall("Variables/Variable[@Class='Parameter']")}
for e in [x for x in experiments if x.get('ActiveObjectClassId')==ped.findtext('Id')]:
 assert {p.findtext('ParameterName') for p in e.findall('Parameters/Parameter')}==set(ped_param_ids.values()), e.findtext('Name')
pvs={e.findtext('Name'):e for e in r.findall('Model/Experiments/ParamVariationExperiment')}
assert set(pvs)=={'PeatonalCorridasApareadas','PeatonalCorridasDemo'}
for name,(runs,modo,archivo,final) in {'PeatonalCorridasApareadas':('60','true','"corridas_peatonales.csv"','18000'),
                                       'PeatonalCorridasDemo':('6','false','"corridas_peatonales_demo.csv"','900')}.items():
 e=pvs[name]
 assert e.get('ActiveObjectClassId')==ped.findtext('Id')
 assert e.findtext('UseFreeformParameters')=='true' and e.findtext('NumberOfRuns')==runs
 assert e.findtext('AllowParallelEvaluations')=='false'
 assert float(e.findtext('ModelTimeProperties/FinalTime'))<=18000, 'PLE limita la Pedestrian Library a 5 h de modelo'
 free={ped_param_ids[x.findtext('Id')]:x.findtext('Expression/Code') for x in e.findall('FreeformParamValue')}
 assert set(free)==set(ped_param_ids.values())
 assert free['semillaPed']=='20260923L + index / 2' and free['molinetesOperativosPed']=='index % 2 == 0 ? 20 : 28'
 assert free['modoFranjaPed']==modo and free['archivoSalidaPed']==archivo
 assert all(free[k] is None for k in ('tamanoTandaPed','intervaloTandaPedSeg','primeraTandaPedSeg','servicioPedSeg'))
# Par i -> semilla 20260923+i, E0 en índice par y E1 en el siguiente.
assert [(20260923+i//2, 20 if i%2==0 else 28) for i in range(4)]==[(20260923,20),(20260923,28),(20260924,20),(20260924,28)]
pasajero_variable_names={x.findtext('Name') for x in agents['Pasajero'].findall('Variables/Variable')}
assert {'tIngresoSistemaPed','enPicoPed','tInicioServicioPed','nombreMolinetePed','servicioAsignadoPed'} <= pasajero_variable_names
libs={x.findtext('LibraryName') for x in r.findall('Model/RequiredLibraryReference')}
assert 'com.anylogic.libraries.pedestrian' in libs
print('OK: XML, unique IDs, clean presentations, pedestrian KPIs, 28 service points and 20/28 experiments')
# The CSV schema, the workbook and the loader share one field order.
import sys, shutil
sys.path.insert(0, str(p.parent))
import cargar_corridas as cc
import openpyxl
book=openpyxl.load_workbook(cc.PLANILLA_OFICIAL)
cc.verificar_encabezados(book[cc.HOJA])
assert [book[cc.HOJA].cell(row,2).value for row in range(8,38)]==list(range(20260923,20260953))
fila_ped=next(f for f in ped.findall('Functions/Function') if f.findtext('Name')=='filaResultadoPed').findtext('Body')
assert fila_ped.count('Integer.toString')+fila_ped.count('String.format')+fila_ped.count('Long.toString')+2==2+len(cc.CAMPOS), 'escenario + "0" desviados + 15 conversiones'
import contextlib, io
with tempfile.TemporaryDirectory(prefix='subte-carga-') as tmp, contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
 tmp=Path(tmp); log=tmp/'consola.txt'; copia=tmp/'copia.xlsx'
 rows=[]
 for i in range(30):
  for esc,esp in (('E0',40+i),('E1',20+i)):
   rows.append(f'CSV_PEATONAL_DEMO;{esc};{20260923+i};480;0;470;480;{esp}.5;{esp+10};0.25;3.5;60;120.0;0.4;0.05;0;0;0')
 log.write_text('ruido de consola\n'+'\n'.join(rows)+'\n')
 assert cc.main([str(log),'--demo'])==0
 assert cc.main([str(log),'--demo','--escribir'])==1, 'demo no puede escribir la planilla oficial'
 assert cc.main([str(log)])==1, 'demo no puede cargarse como producción'
 assert cc.main([str(log),'--demo','--escribir','--salida',str(copia)])==0
 hoja=openpyxl.load_workbook(copia)[cc.HOJA]
 assert hoja['C8'].value==480 and hoja['O8'].value==40.5 and hoja['P8'].value==20.5 and hoja['Q8'].value=='=IF(OR(O8="",P8=""),"",P8-O8)'
 assert hoja['AS37'].value==0 and hoja['AG37'].value==0.4
 (tmp/'incompleta.txt').write_text(rows[0]+'\nCSV_PEATONAL_INCOMPLETO;E1;20260923;480;0;470;479\n')
 assert cc.main([str(tmp/'incompleta.txt'),'--demo'])==1
 (tmp/'impar.txt').write_text(rows[0]+'\n')
 assert cc.main([str(tmp/'impar.txt'),'--demo'])==1
 (tmp/'desbalance.txt').write_text(rows[0]+'\n'+rows[1].replace(';480;0;470;480;',';481;0;470;481;')+'\n')
 assert cc.main([str(tmp/'desbalance.txt'),'--demo'])==1
assert openpyxl.load_workbook(cc.PLANILLA_OFICIAL)[cc.HOJA]['C8'].value is None, 'la planilla oficial debe seguir vacía'
print('OK: CSV_PEATONAL order matches 04-resultados-corridas.xlsx and the loader rejects demo, incomplete or unpaired runs')
